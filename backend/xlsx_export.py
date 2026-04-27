"""Helpers for XLSX exports."""
from io import BytesIO

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_response(filename, sheet_name, headers, rows, percent_columns=None):
    percent_columns = set(percent_columns or [])
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="001F89")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        ws.append(row)

    for col_idx in percent_columns:
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, col_idx).number_format = "0%"

    for column in ws.columns:
        max_len = 0
        col_letter = get_column_letter(column[0].column)
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 2, 62))

    ws.freeze_panes = "A2"
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
