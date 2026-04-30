"""API para indicadores MEL de proyecto con edicion exclusiva para admins."""
from datetime import datetime
import base64
import hashlib
import hmac
import json
import os
import re
import sys
import unicodedata
from typing import Optional

from fastapi import APIRouter, Depends, Header, HTTPException
from pydantic import BaseModel
from sqlalchemy import inspect, text as sql_text
from sqlalchemy.orm import Session

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from database import get_db
from models.mel_proyecto import MelProyectoIndicador
from xlsx_export import xlsx_response

router = APIRouter()

SECRET = os.environ.get("MEL_PROYECTO_SECRET", "mel-proyecto-local-secret")
ADMIN_USER = os.environ.get("MEL_PROYECTO_ADMIN_USER", "admin")
ADMIN_PASSWORD = os.environ.get("MEL_PROYECTO_ADMIN_PASSWORD", "AdminPARES2026!")
SOURCE_SHEET = "Sheet1"
SOURCE_SCHEMA_VERSION = "mel-proyecto-v6-sheet1"
SOURCE_HASH_FIELDS = [
    "source_id", "excel_row", "tipo", "nivel", "indicador", "herramienta",
    "evidencia", "linea_base", "meta", "avance", "porcentaje",
    "fuente_informacion", "frecuencia", "lq", "notas", "estado_fuente",
]


def _workbook_path():
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    candidates = [
        os.environ.get("MEL_PROYECTO_XLSX"),
        os.path.join(root, "MEL PROPOSAL V6.xlsx"),
        os.path.join(root, "mel-proyecto.xlsx"),
        "/app/mel-proyecto.xlsx",
    ]
    return next((path for path in candidates if path and os.path.exists(path)), None)


def _text(value):
    if value is None:
        return None
    clean = str(value).strip()
    return clean or None


def _float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def _cell_text(cell):
    value = cell.value
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _progress_from_porcentaje(value):
    parsed = _float(value)
    if parsed is None:
        return 0.0
    return parsed


def _percent_label(value):
    return f"{round((value or 0) * 100)}%"


def _slug(value):
    normalized = unicodedata.normalize("NFKD", value or "")
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")[:120]


def _status(*values):
    text = " ".join(filter(None, values)).lower()
    return "incompleto" if "completar unep" in text else "completo"


def _record_hash(record):
    payload = json.dumps(
        {field: record.get(field) for field in SOURCE_HASH_FIELDS},
        ensure_ascii=False,
        sort_keys=True,
        default=str,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _ensure_mel_proyecto_schema(db):
    bind = db.get_bind()
    inspector = inspect(bind)
    if "mel_proyecto_indicadores" not in inspector.get_table_names():
        return
    columns = {column["name"] for column in inspector.get_columns("mel_proyecto_indicadores")}
    additions = {
        "evidencia": "TEXT",
        "avance": "TEXT",
        "porcentaje": "TEXT",
        "source_hash": "VARCHAR(64)",
    }
    for name, ddl_type in additions.items():
        if name not in columns:
            db.execute(sql_text(f"ALTER TABLE mel_proyecto_indicadores ADD COLUMN {name} {ddl_type}"))
    db.commit()


def _sign(payload):
    return hmac.new(SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()


def _make_token(username):
    payload = f"{username}|admin"
    raw = f"{payload}|{_sign(payload)}"
    return base64.urlsafe_b64encode(raw.encode()).decode()


def _read_admin_token(authorization):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Credenciales de administrador requeridas")
    try:
        raw = base64.urlsafe_b64decode(authorization.split(" ", 1)[1].encode()).decode()
        username, role, signature = raw.rsplit("|", 2)
    except Exception as exc:
        raise HTTPException(401, "Token invalido") from exc
    payload = f"{username}|{role}"
    if role != "admin" or not hmac.compare_digest(signature, _sign(payload)):
        raise HTTPException(403, "Solo administradores pueden editar MEL Proyecto")
    return {"username": username, "role": role}


def _merge_text(current, incoming):
    incoming = _text(incoming)
    if not incoming:
        return current
    if not current:
        return incoming
    if incoming in current:
        return current
    return f"{current}\n\n{incoming}"


def build_mel_proyecto_records(xlsx_path):
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True)
    if SOURCE_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook does not contain required sheet: {SOURCE_SHEET}")
    ws = wb[SOURCE_SHEET]
    current_tipo = "outcome"
    records = []

    for row_num in range(1, ws.max_row + 1):
        first = _text(ws.cell(row_num, 1).value)
        second = _text(ws.cell(row_num, 2).value)
        if not first and not second:
            continue
        if first == "Indicadores de Outcome":
            current_tipo = "outcome"
            continue
        if first == "Indicadores de Output":
            current_tipo = "output"
            continue
        if first in ("Nivel Outcome", "Output") or second == "Indicador":
            continue
        if not second:
            continue

        indicador = second
        source_id = f"{SOURCE_SHEET.lower()}-{current_tipo}-{row_num:04d}-{_slug(indicador)}"
        herramienta = _text(ws.cell(row_num, 3).value)
        evidencia = _text(ws.cell(row_num, 4).value)
        linea_base_cell = ws.cell(row_num, 5)
        meta_cell = ws.cell(row_num, 6)
        avance_cell = ws.cell(row_num, 7)
        porcentaje_cell = ws.cell(row_num, 8)
        fuente = _text(ws.cell(row_num, 9).value)
        frecuencia = _text(ws.cell(row_num, 10).value)
        lq = _text(ws.cell(row_num, 11).value)
        notas = _text(ws.cell(row_num, 12).value)

        item = {
            "source_id": source_id,
            "excel_row": row_num,
            "tipo": current_tipo,
            "nivel": first,
            "indicador": indicador,
            "herramienta": herramienta,
            "evidencia": evidencia,
            "linea_base": _cell_text(linea_base_cell),
            "meta": _cell_text(meta_cell),
            "meta_numerica": _float(meta_cell.value),
            "avance": _cell_text(avance_cell),
            "valor_actual": _float(avance_cell.value),
            "porcentaje": _cell_text(porcentaje_cell),
            "porcentaje_avance": _progress_from_porcentaje(porcentaje_cell.value),
            "fuente_informacion": fuente,
            "frecuencia": frecuencia,
            "lq": lq,
            "notas": notas,
            "estado_fuente": _status(indicador, herramienta, evidencia, fuente, frecuencia),
            "updated_by": f"seed:{SOURCE_SCHEMA_VERSION}",
        }
        item["source_hash"] = _record_hash(item)
        records.append(item)

    return records


def import_mel_proyecto_xlsx(db, xlsx_path, reset=False):
    _ensure_mel_proyecto_schema(db)
    records = build_mel_proyecto_records(xlsx_path)
    if reset:
        db.query(MelProyectoIndicador).delete()
        db.commit()
    existing = db.query(MelProyectoIndicador).all()
    if existing:
        incoming_ids = {item["source_id"] for item in records}
        incoming_hashes = {item["source_hash"] for item in records}
        existing_ids = {row.source_id for row in existing}
        existing_hashes = {row.source_hash for row in existing if row.source_hash}
        if existing_ids == incoming_ids and existing_hashes == incoming_hashes:
            return 0
        db.query(MelProyectoIndicador).delete()
        db.commit()
    db.add_all(MelProyectoIndicador(**item) for item in records)
    db.commit()
    return len(records)


def seed_mel_proyecto(db):
    _ensure_mel_proyecto_schema(db)
    xlsx_path = _workbook_path()
    if not xlsx_path:
        return
    import_mel_proyecto_xlsx(db, xlsx_path, reset=False)


class LoginIn(BaseModel):
    username: str
    password: str


class MelProyectoPatch(BaseModel):
    tipo: Optional[str] = None
    nivel: Optional[str] = None
    indicador: Optional[str] = None
    herramienta: Optional[str] = None
    evidencia: Optional[str] = None
    linea_base: Optional[str] = None
    meta: Optional[str] = None
    meta_numerica: Optional[float] = None
    avance: Optional[str] = None
    valor_actual: Optional[float] = None
    porcentaje: Optional[str] = None
    porcentaje_avance: Optional[float] = None
    fuente_informacion: Optional[str] = None
    frecuencia: Optional[str] = None
    lq: Optional[str] = None
    notas: Optional[str] = None
    estado_fuente: Optional[str] = None


def _out(row):
    return {
        "id": row.id,
        "source_id": row.source_id,
        "row": row.excel_row,
        "tipo": row.tipo,
        "nivel": row.nivel,
        "indicador": row.indicador,
        "herramienta": row.herramienta,
        "evidencia": row.evidencia,
        "linea_base": row.linea_base,
        "meta": row.meta,
        "meta_numerica": row.meta_numerica,
        "avance": row.avance,
        "valor_actual": row.valor_actual or 0,
        "porcentaje": row.porcentaje,
        "porcentaje_avance": row.porcentaje_avance or 0,
        "fuente_informacion": row.fuente_informacion,
        "frecuencia": row.frecuencia,
        "lq": row.lq,
        "notas": row.notas,
        "estado_fuente": row.estado_fuente,
        "updated_by": row.updated_by,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


def _filtered_query(db, tipo=None, search=None):
    q = db.query(MelProyectoIndicador)
    if tipo and tipo != "all":
        q = q.filter(MelProyectoIndicador.tipo == tipo)
    if search:
        like = f"%{search}%"
        q = q.filter(
            (MelProyectoIndicador.indicador.ilike(like)) |
            (MelProyectoIndicador.nivel.ilike(like)) |
            (MelProyectoIndicador.lq.ilike(like)) |
            (MelProyectoIndicador.notas.ilike(like))
        )
    return q.order_by(MelProyectoIndicador.tipo, MelProyectoIndicador.excel_row)


@router.post("/mel-proyecto/login")
def login(data: LoginIn):
    username = data.username.strip().lower()
    if username != ADMIN_USER.lower() or not hmac.compare_digest(data.password, ADMIN_PASSWORD):
        raise HTTPException(401, "Solo administradores pueden ingresar")
    return {"token": _make_token(username), "username": username, "role": "admin"}


@router.get("/mel-proyecto")
def listar(tipo: Optional[str] = None, search: Optional[str] = None, db: Session = Depends(get_db)):
    seed_mel_proyecto(db)
    rows = _filtered_query(db, tipo, search).all()
    return [_out(row) for row in rows]


@router.get("/mel-proyecto/export")
def exportar_xlsx(tipo: Optional[str] = None, search: Optional[str] = None, db: Session = Depends(get_db)):
    seed_mel_proyecto(db)
    rows = _filtered_query(db, tipo, search).all()
    headers = [
        "Fila Excel", "Tipo", "Nivel / Output", "Indicador", "Instrumento de medicion",
        "Evidencia", "Linea base", "Meta", "Meta numerica", "Avance fuente",
        "Valor actual", "Porcentaje fuente", "Porcentaje avance",
        "Fuente de informacion", "Frecuencia", "LQ", "Notas", "Estado fuente",
        "Actualizado por", "Actualizado en",
    ]
    data = [
        [
            row.excel_row, row.tipo, row.nivel, row.indicador, row.herramienta,
            row.evidencia, row.linea_base, row.meta, row.meta_numerica, row.avance,
            row.valor_actual, row.porcentaje, row.porcentaje_avance,
            row.fuente_informacion, row.frecuencia, row.lq, row.notas, row.estado_fuente,
            row.updated_by, row.updated_at.isoformat() if row.updated_at else None,
        ]
        for row in rows
    ]
    return xlsx_response("mel-proyecto.xlsx", "MEL Proyecto", headers, data, percent_columns={13})


@router.get("/mel-proyecto/resumen")
def resumen(db: Session = Depends(get_db)):
    seed_mel_proyecto(db)
    rows = db.query(MelProyectoIndicador).all()
    outcomes = [row for row in rows if row.tipo == "outcome"]
    outputs = [row for row in rows if row.tipo == "output"]
    avg = sum((row.porcentaje_avance or 0) for row in rows) / len(rows) if rows else 0
    return {
        "total_indicadores": len(rows),
        "outcomes": len(outcomes),
        "outputs": len(outputs),
        "cumplidos": sum(1 for row in rows if (row.porcentaje_avance or 0) >= 1),
        "incompletos": sum(1 for row in rows if row.estado_fuente == "incompleto"),
        "avance_promedio": avg,
    }


@router.patch("/mel-proyecto/{id}")
def actualizar(id: int, data: MelProyectoPatch, authorization: Optional[str] = Header(None), db: Session = Depends(get_db)):
    seed_mel_proyecto(db)
    auth = _read_admin_token(authorization)
    row = db.query(MelProyectoIndicador).get(id)
    if not row:
        raise HTTPException(404, "Indicador no encontrado")
    payload = data.model_dump(exclude_unset=True)
    for key, value in payload.items():
        setattr(row, key, value)
    if "porcentaje_avance" in payload and row.porcentaje_avance is not None:
        row.porcentaje_avance = max(0, min(1.5, row.porcentaje_avance))
    row.estado_fuente = _status(row.indicador, row.herramienta, row.evidencia, row.fuente_informacion, row.frecuencia)
    row.updated_by = auth["username"]
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _out(row)
