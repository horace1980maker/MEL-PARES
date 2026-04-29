"""API para la linea de tiempo editable de Ruta del Proyecto."""
from datetime import datetime
import base64
import hashlib
import hmac
import os
import re
import sys
import unicodedata
from typing import Optional

from fastapi import APIRouter, Depends, Header, HTTPException
from pydantic import BaseModel
from sqlalchemy.orm import Session

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from database import get_db
from models.ruta_timeline import RutaTimelineItem

router = APIRouter()

SECRET = os.environ.get("RUTA_SECRET", "ruta-local-secret")
ADMIN_USER = os.environ.get("RUTA_ADMIN_USER", os.environ.get("MEL_PROYECTO_ADMIN_USER", "admin"))
ADMIN_PASSWORD = os.environ.get("RUTA_ADMIN_PASSWORD", os.environ.get("MEL_PROYECTO_ADMIN_PASSWORD", "AdminPARES2026!"))


def _text(value):
    if value is None:
        return None
    clean = str(value).strip()
    return clean or None


def _int(value):
    if value is None or value == "":
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    match = re.search(r"\d+", str(value))
    return int(match.group(0)) if match else None


def _slug(value):
    normalized = unicodedata.normalize("NFKD", value or "")
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")[:120]


def _norm(value):
    normalized = unicodedata.normalize("NFKD", value or "")
    return normalized.encode("ascii", "ignore").decode("ascii").lower().strip()


def _source_id(row):
    code = _slug(row.get("codigo") or "")
    order = row.get("orden") or row.get("excel_row")
    return f"timeline-{order}-{code}"


def _sign(payload):
    return hmac.new(SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()


def _make_token(username):
    payload = f"{username}|admin"
    raw = f"{payload}|{_sign(payload)}"
    return base64.urlsafe_b64encode(raw.encode()).decode()


def _read_admin_token(authorization):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Credenciales de administrador requeridas")
    try:
        raw = base64.urlsafe_b64decode(authorization.split(" ", 1)[1].encode()).decode()
        username, role, signature = raw.rsplit("|", 2)
    except Exception as exc:
        raise HTTPException(401, "Token invalido") from exc
    payload = f"{username}|{role}"
    if role != "admin" or not hmac.compare_digest(signature, _sign(payload)):
        raise HTTPException(403, "Solo administradores pueden editar Ruta")
    return {"username": username, "role": role}


def build_ruta_timeline_records(xlsx_path):
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True)
    if "TIMELINE" not in wb.sheetnames:
        return []
    ws = wb["TIMELINE"]
    rows = []
    for row_num in range(1, ws.max_row + 1):
        codigo = _text(ws.cell(row_num, 1).value)
        output = _text(ws.cell(row_num, 2).value)
        orden = _int(ws.cell(row_num, 3).value)
        anio = _int(ws.cell(row_num, 4).value)
        mes = _text(ws.cell(row_num, 5).value)
        entregable = _text(ws.cell(row_num, 6).value)
        estado = _text(ws.cell(row_num, 7).value)
        if _norm(codigo) in {"no. entregable", "no entregable"} or _norm(entregable) == "entregable":
            continue
        if not any([codigo, output, orden, anio, mes, entregable, estado]):
            continue
        item = {
            "excel_row": row_num,
            "codigo": codigo or f"fila-{row_num}",
            "output": output,
            "orden": orden,
            "anio": anio,
            "mes": mes,
            "entregable": entregable or "Sin entregable",
            "estado": estado,
            "updated_by": "seed:timeline",
        }
        item["source_id"] = _source_id(item)
        rows.append(item)
    return rows


def import_ruta_timeline_xlsx(db, xlsx_path, reset=False):
    if reset:
        db.query(RutaTimelineItem).delete()
        db.commit()
    if db.query(RutaTimelineItem).count() > 0:
        return 0
    records = build_ruta_timeline_records(xlsx_path)
    db.add_all(RutaTimelineItem(**item) for item in records)
    db.commit()
    return len(records)


def seed_ruta_timeline(db):
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    xlsx_path = os.path.join(root, "MEL PROPOSAL V6.xlsx")
    if not os.path.exists(xlsx_path):
        return
    import_ruta_timeline_xlsx(db, xlsx_path, reset=False)


class LoginIn(BaseModel):
    username: str
    password: str


class RutaTimelinePatch(BaseModel):
    codigo: Optional[str] = None
    output: Optional[str] = None
    orden: Optional[int] = None
    anio: Optional[int] = None
    mes: Optional[str] = None
    entregable: Optional[str] = None
    estado: Optional[str] = None


def _out(row):
    return {
        "id": row.id,
        "source_id": row.source_id,
        "row": row.excel_row,
        "codigo": row.codigo,
        "output": row.output,
        "orden": row.orden,
        "anio": row.anio,
        "mes": row.mes,
        "entregable": row.entregable,
        "estado": row.estado,
        "updated_by": row.updated_by,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


@router.post("/ruta/login")
def login(data: LoginIn):
    username = data.username.strip().lower()
    if username != ADMIN_USER.lower() or not hmac.compare_digest(data.password, ADMIN_PASSWORD):
        raise HTTPException(401, "Solo administradores pueden ingresar")
    return {"token": _make_token(username), "username": username, "role": "admin"}


@router.get("/ruta/timeline")
def listar(db: Session = Depends(get_db)):
    seed_ruta_timeline(db)
    rows = db.query(RutaTimelineItem).order_by(RutaTimelineItem.anio, RutaTimelineItem.orden).all()
    return [_out(row) for row in rows]


@router.patch("/ruta/timeline/{id}")
def actualizar(id: int, data: RutaTimelinePatch, authorization: Optional[str] = Header(None), db: Session = Depends(get_db)):
    seed_ruta_timeline(db)
    auth = _read_admin_token(authorization)
    row = db.query(RutaTimelineItem).get(id)
    if not row:
        raise HTTPException(404, "Entregable no encontrado")
    payload = data.model_dump(exclude_unset=True)
    for key, value in payload.items():
        if isinstance(value, str):
            value = value.strip()
        setattr(row, key, value)
    if not row.codigo:
        raise HTTPException(422, "No. entregable es requerido")
    if not row.entregable:
        raise HTTPException(422, "Entregable es requerido")
    row.updated_by = auth["username"]
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _out(row)
