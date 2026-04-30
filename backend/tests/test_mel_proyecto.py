import os
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[2]
BACKEND_DIR = ROOT / "backend"
DB_FD, DB_PATH = tempfile.mkstemp(suffix=".db")
os.close(DB_FD)
os.environ["DATABASE_URL"] = f"sqlite:///{Path(DB_PATH).as_posix()}"
sys.path.insert(0, str(BACKEND_DIR))

from fastapi.testclient import TestClient  # noqa: E402
from database import engine  # noqa: E402
from main import app  # noqa: E402
from routes.mel_proyecto import build_mel_proyecto_records  # noqa: E402


class MelProyectoTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        cls.client.close()
        engine.dispose()
        try:
            os.remove(DB_PATH)
        except OSError:
            pass

    def test_seed_summary_and_incomplete_indicator(self):
        items = self.client.get("/api/mel-proyecto").json()
        summary = self.client.get("/api/mel-proyecto/resumen").json()
        expected = build_mel_proyecto_records(ROOT / "MEL PROPOSAL V6.xlsx")
        expected_outcomes = sum(1 for row in expected if row["tipo"] == "outcome")
        expected_outputs = sum(1 for row in expected if row["tipo"] == "output")
        expected_incomplete = sum(1 for row in expected if row["estado_fuente"] == "incompleto")

        self.assertEqual(len(items), len(expected))
        self.assertEqual(summary["total_indicadores"], len(expected))
        self.assertEqual(summary["outcomes"], expected_outcomes)
        self.assertEqual(summary["outputs"], expected_outputs)
        self.assertEqual(summary["incompletos"], expected_incomplete)
        self.assertTrue(any(item["estado_fuente"] == "incompleto" for item in items))
        self.assertIn("herramienta", items[0])
        self.assertIn("evidencia", items[0])
        self.assertIn("avance", items[0])
        self.assertIn("porcentaje", items[0])
        self.assertIn("fuente_informacion", items[0])
        self.assertIn("notas", items[0])
        source_row = next(row for row in expected if row["avance"] == "Horacio completa")
        api_row = next(item for item in items if item["source_id"] == source_row["source_id"])
        self.assertEqual(api_row["evidencia"], source_row["evidencia"])
        self.assertEqual(api_row["avance"], source_row["avance"])
        self.assertEqual(api_row["porcentaje"], source_row["porcentaje"])
        self.assertEqual(api_row["porcentaje_avance"], source_row["porcentaje_avance"])
        self.assertEqual(source_row["porcentaje"], "#VALUE!")
        self.assertEqual(source_row["porcentaje_avance"], 0)

    def test_parser_uses_sheet1_and_ignores_other_sheets(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            temp_path = Path(tmp.name)
        try:
            wb = Workbook()
            timeline = wb.active
            timeline.title = "TIMELINE"
            timeline.append(["Output", "TIMELINE should not import"])
            sheet1 = wb.create_sheet("Sheet1")
            sheet1.append(["Indicadores de Outcome"])
            sheet1.append([
                "Nivel Outcome", "Indicador", "Instrumento de medición de mi indicador",
                "Evidencia", "Línea base", "Meta", "Avance", "Porcentaje",
                "Fuente de Información", "Frecuencia", "Learning question ralacionada", "Notas",
            ])
            sheet1.append([
                "Outcome test", "Sheet1 indicator", "Instrumento", "Evidencia",
                0, 10, 5, 0.5, "Fuente", "Mensual", "LQ1", "Nota",
            ])
            sheet3 = wb.create_sheet("Sheet3")
            sheet3.append(["Output", "Sheet3 should not import"])
            wb.active = 0
            wb.save(temp_path)

            records = build_mel_proyecto_records(temp_path)
            self.assertEqual(len(records), 1)
            self.assertTrue(records[0]["source_id"].startswith("sheet1-outcome-"))
            self.assertEqual(records[0]["indicador"], "Sheet1 indicator")
            self.assertEqual(records[0]["porcentaje_avance"], 0.5)
        finally:
            try:
                temp_path.unlink()
            except OSError:
                pass

    def test_admin_only_login_and_update(self):
        items = self.client.get("/api/mel-proyecto").json()
        indicator_id = items[0]["id"]

        bad_login = self.client.post(
            "/api/mel-proyecto/login",
            json={"username": "adel", "password": "bad"},
        )
        self.assertEqual(bad_login.status_code, 401)

        no_auth = self.client.patch(
            f"/api/mel-proyecto/{indicator_id}",
            json={"porcentaje_avance": 0.33},
        )
        self.assertEqual(no_auth.status_code, 401)

        login = self.client.post(
            "/api/mel-proyecto/login",
            json={"username": "admin", "password": "AdminPARES2026!"},
        )
        self.assertEqual(login.status_code, 200)
        self.assertEqual(login.json()["role"], "admin")

        token = login.json()["token"]
        update = self.client.patch(
            f"/api/mel-proyecto/{indicator_id}",
            json={"porcentaje_avance": 0.33, "notas": "Nota de prueba"},
            headers={"Authorization": f"Bearer {token}"},
        )
        self.assertEqual(update.status_code, 200)
        self.assertEqual(update.json()["porcentaje_avance"], 0.33)
        self.assertEqual(update.json()["notas"], "Nota de prueba")

    def test_filter_and_xlsx_export(self):
        outcomes = self.client.get("/api/mel-proyecto", params={"tipo": "outcome"}).json()
        export = self.client.get("/api/mel-proyecto/export", params={"tipo": "outcome"})
        expected = build_mel_proyecto_records(ROOT / "MEL PROPOSAL V6.xlsx")
        expected_outcomes = sum(1 for row in expected if row["tipo"] == "outcome")

        self.assertEqual(len(outcomes), expected_outcomes)
        self.assertEqual(export.status_code, 200)
        self.assertEqual(
            export.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(export.content[:2], b"PK")
        workbook = load_workbook(BytesIO(export.content), data_only=True)
        headers = [cell.value for cell in workbook.active[1]]
        self.assertIn("Evidencia", headers)
        self.assertIn("Avance fuente", headers)
        self.assertIn("Porcentaje fuente", headers)


if __name__ == "__main__":
    unittest.main()
